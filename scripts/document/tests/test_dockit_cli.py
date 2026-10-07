"""`dockit` 命令行(doc_gui_backend.py 的 ops / run / doctor)回归门。

## 为什么有这份(2026-09-30 立)

DocKit 的 GUI 给人用,`dockit` 给 agent 用,两者必须走同一个 gui_run:同一套选项校验、
外发同意门、产出归因。agent 只能看退出码和 JSON,所以这里钉死三件事:

  1. 退出码语义:0 成功 · 1 业务失败 · 2 用法/校验错误 · 75 同目录忙 ——
     gui-* 的「一律 exit 0」契约不许被顺手改掉(Swift 只认信封)。
  2. 诊断型(bidfinal)与预览型(view)的成败不再按「有无新产出」判:
     门检红门是 verdict=red 而不是「后端返回失败(未给出原因)」;view 报告 HTML 产出且默认不开浏览器。
  3. 外发(scan)在没有明确同意时,于列目录/读内容/启动引擎之前拒绝;同意后发现项结构化返回。
     同意路径只用桩 `claude` + 系统禁网验证,不发送任何文档。
  4. (2026-09-30 独立复核后补)覆盖门:与源同名换后缀的产出(b.md → b.docx)、merged.md
     已存在时默认拒绝(would_overwrite),--yes / 勾选才覆盖;分层目录锁让父子目录的操作互斥;
     后台任务可查询;被 SIGTERM 时引擎进程组一起结束、锁文件清掉。

  5. (2026-10-07)顶层帮助的四样(读/写命令、--json 形状、退出码表、「仅在窗口中」)与登记对得上;
     status 读回已装版本与界面记住的设置;settings 读写的是 App 那两个偏好键,值先校验。
  6. (2026-10-07)「配置与更新…」窗口的几项(config / update)整段转给 App 可执行文件:参数、输出、退出码
     原样过去原样回来;找不到 App、App 是不带这组命令的旧版、被信号终止各有一个码,旧版一律不启动
     (它会把这些词当成普通启动、打开窗口)。这里用 sh 桩当 App 程序;编好的真程序在 mac/tests/test_lifecycle_cli.py。
  7. (2026-10-07 下午)共用层补了 update install 与 config status 的 sync_status,窗口的每一项都有命令:帮助里
     不再有「暂无命令」;update install 的转调另有时限(它要下载、验证、替换),其余命令的时限不变。

全部输入都是本测试现造的临时文件;不碰用户文档、不联网、不打开浏览器。
锁与后台任务写进 DOCKIT_CACHE_DIR 指向的临时目录,不碰 ~/Library/Caches。
偏好读写指到 DOCKIT_DEFAULTS_DOMAIN 的临时 plist,已装版本读 DOCKIT_APP_BUNDLE 的临时包,不碰本人偏好与已装 App。
"""
from __future__ import annotations

import fcntl
import hashlib
import json
import os
import plistlib
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[3]
BACKEND = ROOT / "scripts" / "document" / "doc_gui_backend.py"
WRAPPER = ROOT / "mac" / "bin" / "dockit"
VENV_PY = Path.home() / "Dev" / ".venv" / "bin" / "python"
SANDBOX = Path("/usr/bin/sandbox-exec")
NO_NETWORK = "(version 1) (allow default) (deny network*)"
CACHE = Path(tempfile.mkdtemp(prefix="dockit-cache-"))   # 锁与后台任务的隔离目录


def _env(**extra) -> dict:
    # BROWSER=true:即使回归把 --no-open 弄丢,也只会调到 /usr/bin/true,不会弹浏览器抢焦点
    return {**os.environ, "BROWSER": "/usr/bin/true", "DOCKIT_CACHE_DIR": str(CACHE),
            "DOCKIT_DEFAULTS_DOMAIN": str(CACHE / "prefs.plist"),      # 不读写本人的 DocKit 偏好
            "DOCKIT_APP_BUNDLE": str(CACHE / "DocKit.app"), **extra}   # 不读已装的 DocKit.app


@pytest.fixture(scope="module", autouse=True)
def _drop_cache():
    yield
    shutil.rmtree(CACHE, ignore_errors=True)


def dockit(*args: str, env: dict | None = None, timeout: float = 120) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(BACKEND), *map(str, args)], cwd=ROOT,
                          capture_output=True, text=True, timeout=timeout, env=env or _env())


def as_json(proc: subprocess.CompletedProcess) -> dict:
    value = json.loads(proc.stdout)
    assert isinstance(value, dict) and isinstance(value.get("ok"), bool), proc.stdout
    return value


def md5(path: Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()


@pytest.fixture()
def docx_file(tmp_path: Path) -> Path:
    from docx import Document
    doc = Document()
    doc.add_paragraph('正文:"成章",面积10平方米。')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格:"保留",面积20平方米。'
    path = tmp_path / "合成 样例.docx"
    doc.save(path)
    return path


# ─────────────────────────────── help / ops

@pytest.mark.parametrize("args", [["--help"], ["ops", "--help"], ["run", "--help"], ["doctor", "--help"],
                                  ["status", "--help"], ["settings", "--help"], ["settings", "set", "--help"]])
def test_help_exits_zero(args):
    proc = dockit(*args)
    assert proc.returncode == 0, proc.stderr
    assert "dockit" in proc.stdout


def test_ops_json_is_the_gui_catalog():
    listed = as_json(dockit("ops", "--json"))
    gui = as_json(dockit("gui-ops"))
    assert listed == gui and listed["ok"] and len(listed["ops"]) >= 14
    one = as_json(dockit("ops", "convert", "--json"))
    assert one["op"]["id"] == "convert" and one["op"]["target_required"] is True


def test_ops_text_detail_and_unknown():
    detail = dockit("ops", "clean")
    assert detail.returncode == 0 and "scope.table=1" in detail.stdout and "示例" in detail.stdout
    unknown = dockit("ops", "no-such-op", "--json")
    assert unknown.returncode == 2 and as_json(unknown)["error_code"] == "unknown_op"


# ─────────────────────────────── run:成功与失败退出码

def test_run_clean_reports_output_and_keeps_source(docx_file: Path):
    before = md5(docx_file)
    proc = dockit("run", "clean", "--opt", "scope.table=0", "--json", docx_file)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    payload = as_json(proc)
    assert payload["ok"] and payload["succeeded"] == payload["total"] == 1
    outputs = [Path(p) for p in payload["results"][0]["outputs"]]
    assert outputs and all(p.is_file() and p.parent == docx_file.parent for p in outputs)
    assert md5(docx_file) == before


@pytest.mark.parametrize("args, code", [
    (["run", "nope", "{f}"], "unknown_op"),
    (["run", "clean", "--opt", "nope=1", "{f}"], "unknown_option"),
    (["run", "clean", "--opt", "scope.table=maybe", "{f}"], "bad_option_value"),
    (["run", "clean", "--opt", "no-equals", "{f}"], "bad_opt_form"),
    (["run", "convert", "{f}"], "missing_target"),
    (["run", "convert", "--to", "bogus", "{f}"], "bad_target"),
    (["run", "renum", "--to", "bogus", "{f}"], "bad_target"),   # 以前会产出未改动副本并报成功
    (["run", "clean", "--to", "md", "{f}"], "bad_target"),
    (["run", "formatclone", "{f}"], "missing_required"),
    (["run", "fontunify", "{f}"], "confirm_required"),
    (["run", "clean", "{missing}"], "no_valid_files"),
    (["run", "clean", "--bogus-flag", "{f}"], "usage"),
])
def test_run_validation_errors_exit_2_without_writing(docx_file: Path, args, code):
    missing = docx_file.parent / "不存在.docx"
    argv = [a.format(f=docx_file, missing=missing) for a in args]
    before = sorted(p.name for p in docx_file.parent.iterdir())
    proc = dockit(*argv, "--json")
    assert proc.returncode == 2, proc.stdout + proc.stderr
    payload = as_json(proc)
    assert payload["ok"] is False and payload["error_code"] == code and payload["error"]
    assert sorted(p.name for p in docx_file.parent.iterdir()) == before


def test_gui_run_contract_still_exits_zero(docx_file: Path):
    """GUI 契约:校验失败照样 exit 0,错误只在信封里(Swift 看到非零只会报「后端崩了」)。"""
    proc = dockit("gui-run", "--op", "renum", "--to", "bogus", "--files", docx_file)
    assert proc.returncode == 0
    payload = as_json(proc)
    assert payload["ok"] is False and payload["error_code"] == "bad_target"


def test_missing_input_among_valid_is_business_failure(docx_file: Path):
    proc = dockit("run", "clean", "--json", docx_file, docx_file.parent / "缺.docx")
    payload = as_json(proc)
    assert proc.returncode == 1 and payload["ok"] and payload["skipped_missing"] == ["缺.docx"]
    assert payload["all_ok"] is False, "ok 是请求级;all_ok 才说明是否全部成功"


def test_all_inputs_failing_carries_error_text(tmp_path: Path):
    bad = tmp_path / "损坏.docx"
    bad.write_bytes(b"not a docx\n")
    proc = dockit("gui-run", "--op", "clean", "--files", bad)
    payload = as_json(proc)
    assert payload["ok"] is False and payload["error_code"] == "failed"
    assert "全部处理失败" in payload["error"] and payload["results"][0]["ok"] is False
    assert dockit("run", "clean", bad).returncode == 1


def test_dry_run_danger_op_plans_without_writing(docx_file: Path):
    before = sorted(p.name for p in docx_file.parent.iterdir())
    proc = dockit("run", "stripchrome", "--dry-run", "--json", docx_file)
    payload = as_json(proc)
    assert proc.returncode == 0 and payload["dry_run"] is True and payload["results"][0]["ok"]
    assert sorted(p.name for p in docx_file.parent.iterdir()) == before


# ─────────────────────────────── 诊断型 / 预览型成败

def test_bidfinal_red_gate_is_a_verdict_not_a_crash(docx_file: Path):
    """合成 docx 找不到 _project.yaml → 目录门必红(bid_gate rc=2)。"""
    gui = as_json(dockit("gui-run", "--op", "bidfinal", "--to", "pei", "--files", docx_file))
    assert gui["ok"] is True, "红门曾让信封 ok:false 且无 error,GUI 只剩「未给出原因」"
    row = gui["results"][0]
    assert row["verdict"] == "red" and row["ok"] is False and "红门" in row["message"]
    assert dockit("run", "bidfinal", docx_file).returncode == 1


def test_view_reports_html_and_does_not_open(tmp_path: Path):
    md = tmp_path / "预览.md"
    md.write_text("# 标题\n\n正文。\n", encoding="utf-8")
    proc = dockit("run", "view", "--json", "-v", md)
    payload = as_json(proc)
    assert proc.returncode == 0, payload
    html = Path(payload["results"][0]["outputs"][0])
    assert html.suffix == ".html" and html.is_file() and "正文" in html.read_text(encoding="utf-8")
    assert "不打开浏览器" in payload["log"]
    assert dockit("run", "clean", "--open", md).returncode == 2
    shutil.rmtree(html.parent, ignore_errors=True)


# ─────────────────────────────── 外发同意门

def test_scan_refuses_before_touching_directory(tmp_path: Path):
    (tmp_path / "a.md").write_text("合成内容\n", encoding="utf-8")
    for extra in ([], ["--opt", "privacy.cloud_consent=0"]):
        proc = dockit("run", "scan", *extra, "--json", tmp_path)
        assert proc.returncode == 2 and as_json(proc)["error_code"] == "consent_required"


def test_scan_with_consent_returns_findings_via_stub(tmp_path: Path):
    stub = tmp_path / "stub"
    stub.mkdir()
    (stub / "claude").write_text(
        "#!/bin/sh\n"
        "echo '[{\"word\":\"竞品甲公司\",\"category\":\"organization\",\"severity\":\"high\","
        "\"reason\":\"合成\",\"context\":\"竞品甲公司\"}]'\n", encoding="utf-8")
    (stub / "claude").chmod(0o755)
    docs = tmp_path / "docs"
    docs.mkdir()
    (docs / "t.md").write_text("# 标书\n\n本方案由竞品甲公司编制。\n", encoding="utf-8")
    env = _env(PATH=f"{stub}{os.pathsep}{os.environ.get('PATH', '')}")
    argv = [sys.executable, str(BACKEND), "run", "scan", "--opt", "privacy.cloud_consent=1", "--json", str(docs)]
    if SANDBOX.exists():  # 双保险:桩失效也发不出去
        argv = [str(SANDBOX), "-p", NO_NETWORK, *argv]
    proc = subprocess.run(argv, cwd=ROOT, capture_output=True, text=True, timeout=120, env=env)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    row = as_json(proc)["results"][0]
    assert row["findings"] and row["findings"][0]["word"] == "竞品甲公司"


# ─────────────────────────────── 目录互斥 / doctor / 包装脚本

def test_busy_directory_exits_75(docx_file: Path):
    # 锁路径由后端自己算(子进程里 import,不把 doc_dispatch 的 sys.path/monkeypatch 带进 pytest 进程)
    where = subprocess.run(
        [sys.executable, "-c", "import sys; sys.path.insert(0, sys.argv[1]); import doc_gui_backend as b;"
         "b.LOCK_DIR.mkdir(parents=True, exist_ok=True); print(b.lock_path(sys.argv[2]))",
         str(BACKEND.parent), str(docx_file.parent)], capture_output=True, text=True, timeout=60,
        env=_env())
    lock = Path(where.stdout.strip())
    assert where.returncode == 0 and lock.parent.is_dir(), where.stderr
    with open(lock, "a") as fh:
        fcntl.flock(fh.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        proc = dockit("run", "clean", "--json", docx_file)
    lock.unlink()   # 本测试手工持有的锁文件,不是 DocKit 留下的
    assert proc.returncode == 75 and as_json(proc)["error_code"] == "busy"
    assert not list(docx_file.parent.glob("*_fixed*"))


def test_doctor_reports_missing_dependency_as_exit_1():
    probe = (
        "import sys; sys.path.insert(0, sys.argv[1]); sys.argv = ['dockit', 'doctor', '--json'];"
        "import doc_gui_backend as b; b._which = lambda name: None; sys.exit(b.main())"
    )
    proc = subprocess.run([sys.executable, "-c", probe, str(BACKEND.parent)],
                          capture_output=True, text=True, timeout=120)
    payload = json.loads(proc.stdout)
    assert proc.returncode == 1 and payload["ok"] is False
    assert {c["id"] for c in payload["checks"] if not c["ok"]} >= {"uv", "claude"}
    real = dockit("doctor", "--json")
    assert real.returncode == (0 if as_json(real)["ok"] else 1)


@pytest.mark.skipif(not VENV_PY.exists(), reason="包装脚本固定执行 ~/Dev/.venv(本机布局)")
def test_wrapper_help_and_missing_backend(tmp_path: Path):
    ok = subprocess.run([str(WRAPPER), "--help"], capture_output=True, text=True, timeout=60)
    assert ok.returncode == 0 and "dockit" in ok.stdout
    gone = subprocess.run([str(WRAPPER), "ops"], capture_output=True, text=True, timeout=60,
                          env={**os.environ, "HOME": str(tmp_path)})
    assert gone.returncode == 69 and "后端未连接" in gone.stderr


# ─────────────────────────────── 覆盖门(would_overwrite)

def _md(path: Path, text: str = "# 标题\n\n正文。\n") -> Path:
    path.write_text(text, encoding="utf-8")
    return path


@pytest.mark.parametrize("op, args, victim", [
    ("convert", ["--to", "word"], "b.docx"),   # b.md → b.docx
    ("typeset", [], "b.docx"),                 # b.md → 院模板 b.docx
    ("merge", [], "merged.md"),                # 第一份旁的 merged.md
])
def test_existing_user_file_is_not_overwritten_without_consent(tmp_path: Path, op, args, victim):
    src = _md(tmp_path / "b.md")
    inputs = [src] + ([_md(tmp_path / "c.md", "# 第二\n\n第二份。\n")] if op == "merge" else [])
    own = tmp_path / victim
    own.write_bytes(b"USER OWN FILE, not produced by DocKit\n")
    before = md5(own)
    for extra in ([], ["--dry-run"]):
        proc = dockit("run", op, *args, *extra, "--json", *inputs)
        payload = as_json(proc)
        assert proc.returncode == 2 and payload["error_code"] == "would_overwrite", payload
        assert payload["overwrites"] == [str(own)] and md5(own) == before
    gui = as_json(dockit("gui-run", "--op", op, *args, "--files", *inputs))
    assert gui["ok"] is False and gui["error_code"] == "would_overwrite" and md5(own) == before
    plan = as_json(dockit("run", op, *args, "--dry-run", "--yes", "--json", *inputs))
    assert plan["ok"] and plan["results"][0]["overwrites"] == [str(own)] and md5(own) == before


def test_overwrite_with_yes_or_gui_option(tmp_path: Path):
    src = _md(tmp_path / "b.md")
    own = tmp_path / "b.docx"
    own.write_bytes(b"USER OWN FILE\n")
    proc = dockit("run", "convert", "--to", "word", "--yes", "--json", src)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert as_json(proc)["results"][0]["outputs"] == [str(own)] and own.read_bytes()[:2] == b"PK"
    merged = tmp_path / "merged.md"
    merged.write_text("USER OWN\n", encoding="utf-8")
    gui = as_json(dockit("gui-run", "--op", "merge", "--opt", "output.overwrite=1",
                         "--files", src, _md(tmp_path / "c.md", "# 二\n\n第二份。\n")))
    assert gui["ok"] and gui["all_ok"] and "第二份" in merged.read_text(encoding="utf-8")


def test_engine_that_keeps_existing_output_says_so(docx_file: Path):
    """docx → md 的引擎遇到已有 .md 自己跳过:不拦、不覆盖,但要说清为什么没产出。"""
    own = docx_file.with_suffix(".md")
    own.write_text("USER OWN\n", encoding="utf-8")
    proc = dockit("run", "convert", "--to", "md", "--json", docx_file)
    row = as_json(proc)["results"][0]
    assert proc.returncode == 1 and not row["ok"] and "已存在" in row["message"]
    assert own.read_text(encoding="utf-8") == "USER OWN\n"


# ─────────────────────────────── 重跑与原地改写

def test_renum_rerun_still_reports_output(docx_file: Path):
    """copy2 沿用源 mtime 时,无需重排的文档重跑被报成「未产出」(复核实测)。"""
    for _ in range(2):
        proc = dockit("run", "renum", "--json", docx_file)
        assert proc.returncode == 0, proc.stdout
        assert as_json(proc)["results"][0]["outputs"][0].endswith("_序号修正.docx")


def test_txt_merge_uses_only_selected_files(tmp_path: Path):
    (tmp_path / "1.txt").write_text("1\n2\n", encoding="utf-8")
    (tmp_path / "2.txt").write_text("a\n", encoding="utf-8")
    (tmp_path / "3.txt").write_text("NOT SELECTED\n", encoding="utf-8")
    proc = dockit("run", "merge", "--json", tmp_path / "1.txt", tmp_path / "2.txt")
    assert proc.returncode == 0, proc.stdout
    out = tmp_path / "merged.csv"
    assert as_json(proc)["results"][0]["outputs"] == [str(out)]
    assert out.read_text(encoding="utf-8").splitlines() == ["1,a", "2,"]


def test_fontunify_rerun_keeps_the_original_backup(tmp_path: Path):
    from pptx import Presentation
    deck = tmp_path / "e.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "标题 Title"
    prs.save(deck)
    original = md5(deck)
    first = as_json(dockit("run", "fontunify", "--yes", "--json", deck))["results"][0]
    assert first["modified_in_place"] == [str(deck)] and first["outputs"] == [str(deck) + ".backup"]
    second = as_json(dockit("run", "fontunify", "--yes", "--json", deck))["results"][0]
    assert second["ok"] and second["outputs"] and second["outputs"] != first["outputs"]
    assert md5(Path(str(deck) + ".backup")) == original, "第二次不许用改过的文件顶掉原件备份"


# ─────────────────────────────── 分层目录锁

_HOLD = (
    "import sys; sys.path.insert(0, sys.argv[1]); import doc_gui_backend as b\n"
    "with b._dir_locks([sys.argv[2]]):\n"
    "    print('held', flush=True); sys.stdin.read()\n"
)


def _holding(folder: Path):
    proc = subprocess.Popen([sys.executable, "-c", _HOLD, str(BACKEND.parent), str(folder)],
                            stdin=subprocess.PIPE, stdout=subprocess.PIPE, text=True, env=_env())
    assert proc.stdout.readline().strip() == "held"
    return proc


def test_parent_and_child_folders_are_mutually_exclusive(tmp_path: Path):
    child, sibling = tmp_path / "sub", tmp_path / "sib"
    child.mkdir()
    sibling.mkdir()
    for d in (tmp_path, child, sibling):
        _md(d / "x.md")
    locks_before = set((CACHE / "locks").glob("*.lock"))
    cases = [(tmp_path, child / "x.md", 75),   # 父目录在跑,子目录里的操作让路
             (child, tmp_path / "x.md", 75),   # 子目录在跑,父目录的操作(快照含子目录)让路
             (child, sibling / "x.md", 0)]     # 兄弟目录互不影响
    for held, target, code in cases:
        holder = _holding(held)
        try:
            proc = dockit("run", "clean", "--json", target)
        finally:
            holder.communicate("", timeout=30)
        assert proc.returncode == code, (held, target, proc.stdout)
    leftovers = set((CACHE / "locks").glob("*.lock")) - locks_before
    assert not leftovers, f"锁文件没有随释放删掉: {len(leftovers)}"


# ─────────────────────────────── 后台任务 / 信号

def test_background_job_reports_status_and_result(tmp_path: Path):
    src = _md(tmp_path / "b.md", '# 标题\n\n正文:"成章"。\n')
    started = as_json(dockit("run", "clean", "--background", "--json", src))
    job_id = started["job"]["id"]
    assert started["ok"] and started["job"]["status"] == "running"
    deadline = time.monotonic() + 120
    while True:
        proc = dockit("status", job_id, "--json")
        job = as_json(proc)["job"]
        if job["status"] != "running" or time.monotonic() > deadline:
            break
        time.sleep(0.3)
    assert job["status"] == "done" and job["exit_code"] == 0 and proc.returncode == 0, job
    assert job["result"]["all_ok"] and job["result"]["results"][0]["outputs"]
    listing = as_json(dockit("status", "--json"))
    assert any(j["id"] == job_id for j in listing["jobs"])
    missing = dockit("status", "20000101-000000-abcdef", "--json")
    assert missing.returncode == 2 and as_json(missing)["error_code"] == "unknown_job"
    bad = dockit("run", "convert", "--background", "--json", src)   # 校验错误在前台就报
    assert bad.returncode == 2 and as_json(bad)["error_code"] == "missing_target"


def test_sigterm_stops_engine_group_and_releases_locks(tmp_path: Path):
    src = _md(tmp_path / "b.md")
    pidfile = tmp_path / "engine.pid"
    locks_before = set((CACHE / "locks").glob("*.lock"))
    # 引擎换成一个会睡 60 秒的 sh(先把自己的 pid 写进 pidfile),其余走真实 run 路径
    probe = (
        "import sys; sys.path.insert(0, sys.argv[1]); import doc_gui_backend as b\n"
        "engine = ['/bin/sh', '-c', 'echo $$ > \"$0\"; exec sleep 60', sys.argv[2]]\n"
        "b.dd.do_clean = lambda files, opts=None: b.dd._run(engine, 'slow engine')\n"
        "sys.argv = ['dockit', 'run', 'clean', '--json', sys.argv[3]]\n"
        "sys.exit(b.main())\n"
    )
    proc = subprocess.Popen([sys.executable, "-c", probe, str(BACKEND.parent), str(pidfile), str(src)],
                            cwd=tmp_path, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
                            env=_env())
    deadline = time.monotonic() + 30
    while not (pidfile.exists() and pidfile.read_text().strip()) and time.monotonic() < deadline:
        time.sleep(0.1)
    if not pidfile.exists():
        proc.kill()
        pytest.fail("桩引擎没有启动: " + proc.communicate()[1][-400:])
    engine = int(pidfile.read_text())
    proc.send_signal(signal.SIGTERM)
    out, _ = proc.communicate(timeout=30)
    assert proc.returncode == 128 + signal.SIGTERM
    payload = json.loads(out)
    assert payload["error_code"] == "interrupted" and payload["signal"] == signal.SIGTERM
    with pytest.raises(ProcessLookupError):
        os.kill(engine, 0)
    assert not set((CACHE / "locks").glob("*.lock")) - locks_before, "被中断后锁文件没清掉"


def test_scan_consent_on_folder_without_markdown_sends_nothing(tmp_path: Path):
    argv = [sys.executable, str(BACKEND), "run", "scan", "--opt", "privacy.cloud_consent=1", "--json", str(tmp_path)]
    if SANDBOX.exists():
        argv = [str(SANDBOX), "-p", NO_NETWORK, *argv]
    proc = subprocess.run(argv, cwd=ROOT, capture_output=True, text=True, timeout=120, env=_env())
    row = as_json(proc)["results"][0]
    assert proc.returncode == 0 and row["findings"] == [] and row["md_files"] == 0
    assert "没有 .md" in row["message"]


# ─────────────────────────────── 帮助四样 / 读回 / 界面记住的设置(2026-10-07)

def _section(text: str, title: str) -> str:
    """帮助里以 title 起头的那一节(到下一个空行为止)。"""
    m = re.search(r"(?m)^" + re.escape(title) + r".*\n(?:.+\n?)*", text)
    assert m, f"顶层帮助缺「{title}」一节"
    return m.group(0)


def test_top_level_help_has_read_write_json_exit_and_window_only():
    text = dockit("--help").stdout
    reads, writes = _section(text, "读命令"), _section(text, "写命令")
    for sub in ("ops", "status", "doctor", "settings"):
        assert re.search(rf"(?m)^\s+dockit {sub}\b", reads), f"读命令没列 {sub}"
    assert re.search(r"(?m)^\s+dockit run\b", writes) and re.search(r"(?m)^\s+dockit settings set\b", writes)
    for line in ("config status", "update check"):          # 「配置与更新…」窗口的读两项:行首列出
        assert re.search(rf"(?m)^\s+{line}\s", reads), f"读命令没列 {line}"
    assert re.search(r"(?m)^\s+dockit public status\b", reads) and re.search(r"(?m)^\s+dockit public settings set\b", writes)
    for line in ("config export", "config import", "config sync", "update install"):
        assert re.search(rf"(?m)^\s+{line}\s", writes), f"写命令没列 {line}"
    assert "public update install" in writes
    assert "config export" not in reads, "导出会写出指定的文件,不能列在「不写任何文件」的读命令里"
    shapes = _section(text, "--json 输出形状")
    for sub in ("ops", "run", "status", "doctor", "settings"):
        assert re.search(rf"(?m)^\s+{sub}\s+\{{", shapes), f"--json 形状没写 {sub}"
    assert '"error_code"' in shapes
    assert re.search(r'(?m)^\s+config / update .*"command"', shapes) and '"error":{"code","message"}' in shapes
    assert "sync_status{text,at,from,live}" in shapes and "would_install{from,to}" in shapes
    codes = _section(text, "退出码")
    for code in ("0", "1", "2", "69", "75", "128+N"):
        assert re.search(rf"(?m)^\s+{re.escape(code)}\s", codes), f"退出码表缺 {code}"
    for code in ("app_missing", "app_outdated", "app_failed", "timeout"):
        assert code in codes, f"退出码表没写转调失败的 {code}"
    assert "720 秒" in codes, "update install 的转调时限要写在退出码表里"
    window = _section(text, "仅在窗口中")
    assert "升级到新版" not in window, "升级有命令(update install),不是只能在窗口"
    # 「配置与更新…」窗口的每一项都有命令了:升级到新版 → update install,开关下面那句同步状态 → config status 的 sync_status
    assert "暂无命令" not in text and "静默安装" not in text


def test_every_human_feature_is_listed_as_window_only():
    """登记里每个 human 项都要出现在帮助的「仅在窗口中」;每个 command 的子命令都是真子命令。"""
    import yaml
    spec = yaml.safe_load((ROOT / "mac" / "project.yaml").read_text(encoding="utf-8"))["sop"]["agent_cli"]
    text = dockit("--help").stdout
    window = _section(text, "仅在窗口中")
    features = spec["features"]
    assert features and all(sum(k in f for k in ("command", "human", "missing")) == 1 for f in features)
    absent = [f["name"] for f in features if "human" in f and f["name"] not in window]
    assert not absent, f"帮助的「仅在窗口中」没列: {absent}"
    for row in [spec["readback"], *(f["command"] for f in features if "command" in f)]:
        words = row.split()
        assert words[0] == "dockit" and dockit(*words[1:], "--help").returncode == 0, row
    for f in features:   # 暂缺项在帮助里也要有去向,不能只在登记里
        if "missing" in f:
            assert f["name"].rstrip("…") in text, f["name"]
    by_name = {f["name"]: f for f in features}
    assert by_name["升级到新版"].get("command") == "dockit update install"
    assert by_name["iCloud 配置同步状态"].get("command") == "dockit config status"


def _fake_app(root: Path, version: str = "9.8.7", build: str = "654") -> Path:
    app = root / "DocKit.app"
    (app / "Contents").mkdir(parents=True)
    (app / "Contents" / "Info.plist").write_bytes(plistlib.dumps({
        "CFBundleIdentifier": "cyou.tianli.DocTools", "CFBundleShortVersionString": version, "CFBundleVersion": build}))
    return app


def _tree(root: Path) -> dict:
    return {str(p.relative_to(root)): (p.stat().st_size, p.stat().st_mtime_ns) for p in sorted(root.rglob("*"))}


def test_status_reads_back_version_settings_and_jobs_without_writing(tmp_path: Path):
    app = _fake_app(tmp_path)
    env = _env(DOCKIT_CACHE_DIR=str(tmp_path / "cache"), DOCKIT_DEFAULTS_DOMAIN=str(tmp_path / "prefs.plist"),
               DOCKIT_APP_BUNDLE=str(app))
    before = _tree(tmp_path)
    proc = dockit("status", "--json", env=env)
    payload = as_json(proc)
    assert proc.returncode == 0 and payload["ok"] and payload["jobs"] == []
    assert payload["app"] == {"path": str(app), "installed": True, "bundle_id": "cyou.tianli.DocTools",
                              "version": "9.8.7", "build": "654"}
    assert payload["settings"] == {"last_operation": None, "target_formats": {}}
    assert "9.8.7 (654)" in dockit("status", env=env).stdout
    for args in (["settings", "--json"], ["ops", "--json"], ["ops", "convert"], ["status"]):
        assert dockit(*args, env=env).returncode == 0
    assert _tree(tmp_path) == before, "读命令写了文件"      # 读命令不建缓存目录、不建偏好文件
    gone = as_json(dockit("status", "--json", env={**env, "DOCKIT_APP_BUNDLE": str(tmp_path / "none.app")}))
    assert gone["app"]["installed"] is False and gone["app"]["version"] is None


def test_settings_set_validates_then_reads_back_only_its_two_keys(tmp_path: Path):
    prefs = tmp_path / "prefs.plist"
    # 域里已有界面自己的其他内容:写命令不许动它,读命令不许外带
    prefs.write_bytes(plistlib.dumps({"NSWindow Frame Main": "1 2 3 4", "dockit.targetFormats": {"renum": "tabfig"}}))
    env = _env(DOCKIT_DEFAULTS_DOMAIN=str(prefs))
    first = as_json(dockit("settings", "--json", env=env))
    assert first["settings"] == {"last_operation": None, "target_formats": {"renum": "tabfig"}}
    assert "NSWindow" not in json.dumps(first)
    for args, code in ([["last_operation", "no-such-op"], "unknown_op"],
                       [["target_formats.no-such-op", "md"], "unknown_op"],
                       [["target_formats.convert", "no-such-target"], "bad_target"],
                       [["target_formats.clean", "md"], "bad_target"],      # 规范化没有目标可选
                       [["window_frame", "0 0 1 1"], "unknown_setting"]):
        bad = dockit("settings", "set", *args, "--json", env=env)
        assert bad.returncode == 2 and as_json(bad)["error_code"] == code, (args, bad.stdout)
    assert plistlib.loads(prefs.read_bytes()) == {"NSWindow Frame Main": "1 2 3 4",
                                                  "dockit.targetFormats": {"renum": "tabfig"}}, "校验失败却写了偏好"
    usage = dockit("settings", "set", "last_operation", "--json", env=env)       # 缺值:用法错误也给 JSON
    assert usage.returncode == 2 and as_json(usage)["error_code"] == "usage"
    done = as_json(dockit("settings", "set", "last_operation", "convert", "--json", env=env))
    assert done["changed"] == {"key": "last_operation", "from": None, "to": "convert"}
    done = as_json(dockit("settings", "--json", "set", "target_formats.convert", "md", env=env))   # --json 在前也认
    assert done["changed"] == {"key": "target_formats.convert", "from": None, "to": "md"}
    again = as_json(dockit("settings", "set", "target_formats.convert", "word", "--json", env=env))
    assert again["changed"]["from"] == "md" and again["settings"]["target_formats"] == {"convert": "word", "renum": "tabfig"}
    text = dockit("settings", env=env)
    assert text.returncode == 0 and "convert" in text.stdout and "renum=tabfig" in text.stdout
    assert as_json(dockit("status", "--json", env=env))["settings"] == again["settings"]      # status 读到同一份
    # App 读的就是这两个键(ViewModel.portablePreferenceKeys);别的键原样留着
    assert plistlib.loads(prefs.read_bytes()) == {
        "NSWindow Frame Main": "1 2 3 4", "dockit.lastOperation": "convert",
        "dockit.targetFormats": {"convert": "word", "renum": "tabfig"}}


def test_settings_keys_match_the_app_preference_keys():
    """settings 写的键名必须是界面读的那两个;Swift 侧改名而这里没跟,命令就成了写给空气。"""
    swift = (ROOT / "mac" / "Sources" / "ViewModel.swift").read_text(encoding="utf-8")
    assert 'portablePreferenceKeys = ["dockit.lastOperation", "dockit.targetFormats"]' in swift
    backend = BACKEND.read_text(encoding="utf-8")
    assert 'PREF_LAST_OP, PREF_TARGETS = "dockit.lastOperation", "dockit.targetFormats"' in backend
    assert 'or "cyou.tianli.DocTools"' in backend      # 默认域 = App 的 bundle id


# ─────────────────────────────── 「配置与更新…」窗口的几项:config / update 转调 App 可执行文件

def _stub_app(root: Path, script: str, verbs=("config", "update"), name: str = "Stub.app",
              bundle_id: str = "test.dockit.stub") -> Path:
    """一个假的 App 包:可执行文件是 sh 桩;verbs 为 None 时 Info.plist 不声明命令字(= 旧版)。"""
    app = root / name
    (app / "Contents" / "MacOS").mkdir(parents=True)
    exe = app / "Contents" / "MacOS" / "Stub"
    exe.write_text("#!/bin/sh\n" + script, encoding="utf-8")
    exe.chmod(0o755)
    info = {"CFBundleIdentifier": bundle_id, "CFBundleExecutable": "Stub",
            "CFBundleShortVersionString": "9.8.7", "CFBundleVersion": "654"}
    if verbs is not None:
        info["DocKitCommandVerbs"] = list(verbs)
    (app / "Contents" / "Info.plist").write_bytes(plistlib.dumps(info))
    return app


def test_config_and_update_are_forwarded_whole_and_come_back_unchanged(tmp_path: Path):
    seen = tmp_path / "argv"
    app = _stub_app(tmp_path, f"""printf '%s\\n' "$@" > '{seen}'; pwd >> '{seen}'
printf '{{"ok":false,"command":"config sync","error":{{"code":"confirmation_required","message":"m"}}}}\\n'
printf 'to stderr\\n' >&2
exit 2
""")
    env = _env(DOCKIT_APP_BUNDLE=str(app))
    where = tmp_path / "typed here"
    where.mkdir()
    done = subprocess.run([sys.executable, str(BACKEND), "config", "sync", "on", "-o", "相对 路径.json", "--json"],
                          cwd=where, capture_output=True, text=True, timeout=60, env=env)
    assert done.returncode == 2 and done.stderr == "to stderr\n"
    assert as_json(done)["error"] == {"code": "confirmation_required", "message": "m"}     # 一个字段都没补、没改
    assert seen.read_text(encoding="utf-8").splitlines() == ["config", "sync", "on", "-o", "相对 路径.json", "--json",
                                                              str(where.resolve())]   # 参数原样,目录是敲命令的目录
    update = dockit("update", "check", env=env)
    assert update.returncode == 2 and seen.read_text(encoding="utf-8").splitlines()[:2] == ["update", "check"]
    # 原有命令不经转调:config 只认第一个词
    assert as_json(dockit("ops", "--json", env=env))["ok"] is True
    assert dockit("--json", "config", "status", env=env).returncode == 2


def test_config_without_an_app_or_with_an_old_app_never_starts_anything(tmp_path: Path):
    marker = tmp_path / "started"
    absent = _env(DOCKIT_APP_BUNDLE=str(tmp_path / "Absent.app"))
    missing = dockit("config", "status", "--json", env=absent)
    assert missing.returncode == 1 and as_json(missing) == {
        "ok": False, "command": "config status", "error": as_json(missing)["error"]}
    assert as_json(missing)["error"]["code"] == "app_missing"
    text = dockit("update", "check", env=absent)
    assert text.returncode == 1 and text.stdout == "" and "App 可执行文件" in text.stderr
    # 旧版程序不认 config / update,会当成普通启动打开窗口:没声明命令字的包一律不启动
    for name, verbs in (("Old.app", None), ("Partial.app", ("update",))):
        old = _env(DOCKIT_APP_BUNDLE=str(_stub_app(tmp_path, f"touch '{marker}'\n", verbs, name)))
        refused = dockit("config", "sync", "on", "--yes", "--json", env=old)
        assert refused.returncode == 1 and as_json(refused)["error"]["code"] == "app_outdated", refused.stdout
        assert not marker.exists(), f"{name} 被启动了"
    assert dockit("update", "check", env=old).returncode == 0 and marker.exists()      # 声明了的词照常转调
    # 帮助不需要 App:没装也看得到用法,退出 0
    for words in (["config", "--help"], ["config", "sync", "--help"], ["update", "check", "-h"]):
        shown = dockit(*words, env=absent)
        assert shown.returncode == 0 and "usage: dockit config [status] [--json]" in shown.stdout, words
        assert "退出码与 error.code" in shown.stdout
    # 被信号终止:没有结果可带回,按同一形状说明
    killed = _env(DOCKIT_APP_BUNDLE=str(_stub_app(tmp_path, "kill -9 $$\n", name="Killed.app")))
    failed = dockit("config", "status", "--json", env=killed)
    assert failed.returncode == 1 and as_json(failed)["error"]["code"] == "app_failed"


@pytest.mark.skipif(not VENV_PY.exists(), reason="包装脚本固定执行 ~/Dev/.venv(本机布局)")
def test_wrapper_inside_a_bundle_forwards_to_its_own_bundle(tmp_path: Path):
    """装进包的 bin/dockit 把 config / update 交给自己所在的包;经软链启动(~/.local/bin/dockit)也一样。"""
    seen = tmp_path / "which"
    app = _stub_app(tmp_path, f"printf 'own bundle\\n' > '{seen}'; printf '{{\"ok\":true}}\\n'\n", name="Own.app")
    entry = app / "Contents" / "Resources" / "bin" / "dockit"
    entry.parent.mkdir(parents=True)
    shutil.copy2(WRAPPER, entry)
    link = tmp_path / "dockit"
    link.symlink_to(entry)
    env = {k: v for k, v in _env().items() if k != "DOCKIT_APP_BUNDLE"}
    env["DOCKIT_ENTRY_APP"] = str(tmp_path / "Stale.app")          # 外面带进来的旧值不算数
    for start in (entry, link):
        seen.unlink(missing_ok=True)
        done = subprocess.run([str(start), "config", "status", "--json"], capture_output=True, text=True, timeout=60, env=env)
        assert done.returncode == 0 and as_json(done)["ok"] is True, done.stdout + done.stderr
        assert seen.read_text() == "own bundle\n"
    # 不在包里的入口(工作树的 mac/bin/dockit)不导出,后端用默认装机位置;这里只核它没把旧值带下去
    outside = subprocess.run([str(WRAPPER), "config", "status", "--json"], capture_output=True, text=True, timeout=60,
                             env=dict(env, DOCKIT_APP_BUNDLE=str(tmp_path / "Absent.app")))
    assert outside.returncode == 1 and as_json(outside)["error"]["code"] == "app_missing"


def test_lifecycle_help_lines_are_listed_from_one_place():
    """顶层帮助里共用层那几行与 `dockit config --help` 用的是同一组常量;逐字是否等于编好的程序,
    由 mac/tests/test_lifecycle_cli.py 拿真程序核对。"""
    top = dockit("--help").stdout
    own = dockit("config", "--help").stdout          # 测试环境没有 App:后端自己的那份
    lines = own.splitlines()
    reads = lines[lines.index("读（不写任何文件或状态）:") + 1:lines.index("写:")]
    writes = lines[lines.index("写:") + 1:next(i for i, line in enumerate(lines) if line.startswith("--json："))]
    assert len(reads) == 2 and len(writes) == 4
    assert reads[0].lstrip().startswith("config status") and "同步状态" in reads[0]
    assert writes[3].lstrip().startswith("update install --yes")
    for line in reads + writes:
        assert "\n" + line + "\n" in top, line
    assert "暂无命令" not in own and "sync_status{text, at, from, live}" in own and "update install →" in own
    cli = (ROOT / "mac" / "Sources" / "AppLifecycleCLI.swift").read_text(encoding="utf-8")
    for line in reads + writes:                      # 与共用层源码里的帮助行逐字相同(命令名代入后)
        named = line.replace("dockit config status", "\\(command) config status").replace(
            "dockit update check", "\\(command) update check")
        assert named in cli, line


def test_update_install_is_forwarded_with_its_own_time_limit(tmp_path: Path):
    """update install 要下载、验证、替换(共用层自己最多等 330 + 20 + 330 秒),转调给它另一个时限;其余命令照旧。
    两条分支都用会睡的桩实测:时限在子进程里缩短,不改生产常量。dockit public 走同一个转调。"""
    seen = tmp_path / "argv"
    script = f"""printf '%s\\n' "$@" > '{seen}'
sleep 2
printf '{{"ok":true,"command":"update install","installed":false,"state":"up_to_date"}}\\n'
"""
    own = _stub_app(tmp_path, script)
    public = _stub_app(tmp_path, script, ("status", "settings", "config", "update", "help"), "Public.app", PUBLIC_ID)
    probe = (
        "import sys; sys.path.insert(0, sys.argv[1]); import doc_gui_backend as b\n"
        "assert (b.LIFECYCLE_SECONDS, b.LIFECYCLE_INSTALL_SECONDS) == (60, 720), 'production limits changed'\n"
        "assert b.lifecycle_seconds(['update', '--json', 'install', '--yes']) == 720\n"
        "assert b.lifecycle_seconds(['update', 'check']) == b.lifecycle_seconds(['config', 'import', 'install']) == 60\n"
        "b.LIFECYCLE_SECONDS, b.LIFECYCLE_INSTALL_SECONDS = 1, 30\n"
        "sys.argv = ['dockit', *sys.argv[2:]]\n"
        "sys.exit(b.main())\n"
    )
    env = _env(DOCKIT_APP_BUNDLE=str(own), DOCKIT_PUBLIC_APP=str(public))

    def typed(*words):
        return subprocess.run([sys.executable, "-c", probe, str(BACKEND.parent), *words], capture_output=True,
                              text=True, timeout=60, env=env)

    for prefix in ((), ("public",)):
        done = typed(*prefix, "update", "install", "--yes", "--dry-run", "--json")
        assert done.returncode == 0, done.stdout + done.stderr
        assert as_json(done) == {"ok": True, "command": "update install", "installed": False, "state": "up_to_date"}
        assert seen.read_text(encoding="utf-8").splitlines() == ["update", "install", "--yes", "--dry-run", "--json"]
        cut = typed(*prefix, "update", "check", "--json")              # 其余命令的时限没有跟着放开
        assert cut.returncode == 1 and as_json(cut)["error"]["code"] == "timeout", cut.stdout + cut.stderr
        assert "1 秒内没有结束" in as_json(cut)["error"]["message"]


# ─────────────────────────────── 公开版自己的状态与设置:dockit public 转调公开包的可执行文件

PUBLIC_ID = "io.github.zengtianli.DocTools"


def test_public_words_go_to_the_public_bundle_unchanged(tmp_path: Path):
    seen = tmp_path / "argv"
    script = f"""printf '%s\\n' "$DOCKIT_CLI_NAME" "$@" > '{seen}'
printf '{{"ok":true,"command":"status"}}\\n'; printf 'note\\n' >&2; exit 0
"""
    app = _stub_app(tmp_path, script, ("status", "settings", "config", "update", "help"), "Public.app", PUBLIC_ID)
    env = _env(DOCKIT_PUBLIC_APP=str(app))
    done = dockit("public", "status", "--json", env=env)
    assert done.returncode == 0 and as_json(done) == {"ok": True, "command": "status"} and done.stderr == "note\n"
    assert seen.read_text(encoding="utf-8").splitlines() == ["dockit public", "status", "--json"]   # 命令怎么敲的告诉它
    dockit("public", "settings", "set", "target_formats.convert", "md", "--json", env=env)
    assert seen.read_text(encoding="utf-8").splitlines()[1:] == ["settings", "set", "target_formats.convert", "md", "--json"]
    dockit("public", "config", "sync", "on", "--yes", env=env)
    assert seen.read_text(encoding="utf-8").splitlines()[1:] == ["config", "sync", "on", "--yes"]
    for words in (["public"], ["public", "--help"], ["public", "-h"], ["public", "help"]):     # 装了新版:完整帮助由它自己给
        seen.unlink()
        assert dockit(*words, env=env).returncode == 0 and seen.read_text(encoding="utf-8").splitlines() == ["dockit public", "help"]
    # 同一个测试域的自用版命令不受影响:config 仍交给自用版那个包,不是公开包
    own = dockit("config", "status", "--json", env=env)
    assert own.returncode == 1 and as_json(own)["error"]["code"] == "app_missing"


def test_public_never_starts_a_bundle_that_does_not_declare_the_word(tmp_path: Path):
    marker = tmp_path / "started"
    cases = (("Old.app", None, PUBLIC_ID, "app_outdated"),                   # 现装的旧公开版:不认命令字,会开窗口
             ("Partial.app", ("config", "update"), PUBLIC_ID, "app_outdated"),
             ("Private.app", ("status",), "cyou.tianli.DocTools", "app_missing"),   # 那个位置上不是公开版
             ("Lookalike.app", ("status",), PUBLIC_ID + "Evil", "app_missing"))
    for name, verbs, bundle_id, code in cases:
        env = _env(DOCKIT_PUBLIC_APP=str(_stub_app(tmp_path, f"touch '{marker}'\n", verbs, name, bundle_id)))
        refused = dockit("public", "status", "--json", env=env)
        assert refused.returncode == 1 and as_json(refused)["error"]["code"] == code, (name, refused.stdout)
        assert as_json(refused)["command"] == "status"
        text = dockit("public", "status", env=env)
        assert text.returncode == 1 and text.stdout == "" and text.stderr.startswith("dockit: ")
        for words in (["public"], ["public", "--help"], ["public", "help"]):      # 帮助不启动它,给这一层自己的说明
            shown = dockit(*words, env=env)
            assert shown.returncode == 0 and "usage: dockit public" in shown.stdout and "app_outdated" in shown.stdout
        assert not marker.exists(), f"{name} 被启动了"
    absent = dockit("public", "status", "--json", env=_env(DOCKIT_PUBLIC_APP=str(tmp_path / "Absent.app")))
    assert absent.returncode == 1 and as_json(absent)["error"]["code"] == "app_missing"
    # 测试用的换了 bundle id 的公开包(…DocTools.<后缀>)认
    suffixed = _env(DOCKIT_PUBLIC_APP=str(_stub_app(tmp_path, "printf '{\"ok\":true}\\n'\n", ("status",), "Test.app",
                                                    PUBLIC_ID + ".LifecycleTest")))
    assert as_json(dockit("public", "status", "--json", env=suffixed))["ok"] is True
