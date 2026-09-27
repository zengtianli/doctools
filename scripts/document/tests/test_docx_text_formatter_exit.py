"""Real CLI regression: failed normalization must reach the caller's exit code."""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document


SCRIPT = Path(__file__).resolve().parent.parent / "docx_fmt.py"


def make_good(path: Path) -> None:
    document = Document()
    document.add_paragraph('恢复, "成功" 20平方米')
    document.save(path)


def run_cli(*paths: Path) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(SCRIPT), "text", *map(str, paths)],
        capture_output=True, text=True, timeout=30,
    )


def assert_normalized(source: Path) -> None:
    output = source.with_name(f"{source.stem}_fixed.docx")
    assert output.is_file()
    document = Document(output)
    assert document.paragraphs[0].text == "恢复， “成功” 20m²"


def test_success_exits_zero_and_preserves_source(tmp_path: Path):
    source = tmp_path / "good.docx"
    make_good(source)
    before = source.read_bytes()

    result = run_cli(source)

    assert result.returncode == 0, result.stdout + result.stderr
    assert_normalized(source)
    assert source.read_bytes() == before


def test_corrupt_docx_exits_one_without_output(tmp_path: Path):
    source = tmp_path / "bad.docx"
    source.write_bytes(b"not a DOCX zip archive\n")
    before = source.read_bytes()

    result = run_cli(source)

    assert result.returncode == 1, result.stdout + result.stderr
    assert "处理失败" in result.stdout
    assert source.read_bytes() == before
    assert set(tmp_path.iterdir()) == {source}


@pytest.mark.parametrize("bad_first", [True, False])
def test_mixed_batch_reports_failure_and_finishes_good_file(tmp_path: Path, bad_first: bool):
    good, bad = tmp_path / "good.docx", tmp_path / "bad.docx"
    make_good(good)
    bad.write_bytes(b"not a DOCX zip archive\n")
    before = {path: path.read_bytes() for path in (good, bad)}

    paths = (bad, good) if bad_first else (good, bad)
    result = run_cli(*paths)

    assert result.returncode == 1, result.stdout + result.stderr
    assert_normalized(good)
    assert not bad.with_name("bad_fixed.docx").exists()
    assert {path: path.read_bytes() for path in before} == before
