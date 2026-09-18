"""bid_body.py 回归门 —— 把 2026-09-17 海宁标实际踩过的坑固化成断言。

## 为什么有这份测试

海宁第 6 章被一次性脚本按「doc.paragraphs 列表」重排：表格不是段落，没跟着动，
结果 9 张表 + sectPr 全部落到文首，整份结构损坏；同一轮还把图插进了「引表短行」
和表题之间，并在同一秒内连续写回、后一份备份盖掉了前一份原始备份。

## 判据

  1. 挪图后：图在正文之后；每张表仍紧跟自己的表题；sectPr 仍在末位
  2. 正文末段若是引出表格的短行，图停在它之前
  3. 损坏形态（表在首个标题之前 / sectPr 不在末位）check 必须判红
  4. 同秒两次写回，两份备份都在
  6. 二级标题下平铺的 N）正文条目：h3group 只插三级标题与引导段、组内编号从 1）重排、
     追加条目也是正文段、表仍跟表题；计划条数对不上或重复执行必须拒绝
  7. 条目编号体例：标题只用 1.1.1.1 式编号，标题样式里出现 N）/（十九）check 必须判红；
     正文条目跳号判红，同一标题下几组并列清单各自从 1）起不算断点；h4num 已退役
  5. 来源渲染：简称→正式名、PDF 页→印刷页按已核偏移换算、同文件相邻页合并、
     无法确定的进 errs 而不是猜
"""
from __future__ import annotations

import base64
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HERE))
import bid_body as B  # noqa: E402

TOOL = HERE / "bid_body.py"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def build(tmp_path, lead_line=False):
    png = tmp_path / "x.png"
    png.write_bytes(PNG)
    d = Document()
    d.add_heading("1.1 概况", level=3)
    d.add_picture(str(png))
    d.add_paragraph("图1-1 技术联系")
    d.add_heading("1.1.1 任务定位", level=4)
    d.add_paragraph("第一段正文，说明任务定位与总体认识。")
    d.add_paragraph("第二段正文，说明证据来源与核验方式。")
    if lead_line:
        d.add_paragraph("指标定义与取数判断")
    d.add_paragraph("表1-1 指标一览")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "指标"
    d.add_heading("1.1.2 资源约束", level=4)
    d.add_paragraph("第三段正文。")
    out = tmp_path / "t.docx"
    d.save(out)
    return out


def seq(path):
    doc = B.Doc(path)
    out = []
    for k in doc.body:
        if k.tag == B.w("tbl"):
            out.append("TBL")
        elif k.tag == B.w("sectPr"):
            out.append("SECT")
        elif doc.has_img(k):
            out.append("IMG")
        else:
            out.append(B.ptext(k).strip()[:6])
    return out


def run(*args):
    return subprocess.run([sys.executable, str(TOOL), *map(str, args)], capture_output=True, text=True)


def test_figs_moves_block_after_text_and_keeps_tables(tmp_path):
    f = build(tmp_path)
    assert run("check", f).returncode == 2
    r = run("figs", f, "--apply")
    assert r.returncode == 0, r.stdout + r.stderr
    s = seq(f)
    assert s.index("IMG") > s.index("第二段正文，")
    assert s[s.index("IMG") + 1] == "图1-1 技"
    assert s[s.index("TBL") - 1] == "表1-1 指"          # 表仍紧跟表题
    assert s[-1] == "SECT"
    assert run("check", f).returncode == 0


def test_figs_stops_before_table_lead_line(tmp_path):
    f = build(tmp_path, lead_line=True)
    assert run("figs", f, "--apply").returncode == 0
    s = seq(f)
    i = s.index("IMG")
    assert s[i + 2] == "指标定义与取" and s[i + 3] == "表1-1 指"   # 引表短行没有被图拆开


def test_check_flags_tables_dumped_to_front(tmp_path):
    f = build(tmp_path)
    doc = B.Doc(f)
    tbl = next(k for k in doc.body if k.tag == B.w("tbl"))
    sect = doc.body[-1]
    doc.body.remove(tbl)
    doc.body.remove(sect)
    doc.body.insert(0, sect)
    doc.body.insert(0, tbl)                              # 复现「按段落列表重排」的损坏形态
    doc.save()
    r = run("check", f)
    assert r.returncode == 2
    assert "sectPr 不在 body 末位" in r.stdout and "表格位于首个标题之前" in r.stdout


def test_same_second_backups_do_not_overwrite(tmp_path):
    f = build(tmp_path)
    for _ in range(3):
        B.Doc(f).save()
    assert len(list(tmp_path.glob("t.docx.bak-*"))) == 3


RULES = {
    "aliases": {"水网规划": "《海宁水网建设规划》", "水资源总体规划": "《水资源总体规划》"},
    "pdf_offset": {"《海宁水网建设规划》": 3},
    "overrides": {"S-05，资料包": ""},
    "drop_notes": ["拟议方法"],
}


@pytest.mark.parametrize("detail,want", [
    ("S-05，水网规划PDF第19页、印刷第16页", "（《海宁水网建设规划》第16页）"),
    ("S-05，水网规划PDF第13页", "（《海宁水网建设规划》第10页）"),                     # 仅 PDF 页 → 按已核偏移换算
    ("S-05，水网规划PDF第42页、印刷第39页；PDF第45页、印刷第42页", "（《海宁水网建设规划》第39页、第42页）"),
    ("S-01，招标需求印刷第6页；拟议方法", "（招标文件“招标需求”第6页）"),
    ("S-01，第二章“项目背景”；S-05，资料包", "（招标文件第二章“项目背景”）"),
])
def test_render(detail, want):
    errs = []
    assert B.render(detail, RULES, errs) == want and not errs


@pytest.mark.parametrize("detail", [
    "S-05",                                   # 裸号
    "S-05，某不认识的规划PDF第3页",              # 文件名无法识别
    "S-05，水资源总体规划PDF第8页",              # 只有 PDF 页且该文件偏移未核
    "S-03，案例参考，段617—624",                # 无页码无 override
])
def test_render_refuses_to_guess(detail):
    errs = []
    B.render(detail, RULES, errs)
    assert errs


def build_flat(tmp_path):
    d = Document()
    d.add_heading("7 技术路线", level=1)
    d.add_heading("7.1 现状调查", level=2)
    d.add_heading("7.1.1 思路", level=3)
    d.add_heading("7.1.1.1 样板条目", level=4)
    d.add_paragraph("样板正文" * 20)
    d.add_heading("7.2 实施方案", level=2)
    d.add_paragraph("节首说明" * 20)
    for i, name in enumerate(["方案框架", "规划约束", "供需平衡", "空间布局", "重点工程"], 1):
        d.add_paragraph(f"{i}）{name}")
        d.add_paragraph(f"{name}正文" * 15)
    d.add_paragraph("表7-1 指标")
    d.add_table(rows=2, cols=2)
    out = tmp_path / "flat.docx"
    d.save(out)
    return out


PLAN = """
sections:
  - h2: "7.2 实施方案"
    groups:
      - {title: "7.2.1 总体思路", n: 2, lead: "引导段一（招标文件第二章）。"}
      - title: "7.2.2 布局与工程"
        n: 3
        lead: "引导段二。"
        append:
          - {h4: "保障措施", paras: ["新增正文（招标文件第二章）。"]}
"""


def items_under(doc):
    return [B.ptext(k).strip() for k in doc.body if doc.is_text(k) and B.ITEM.match(B.ptext(k).strip())]


def test_h3group_inserts_levels_and_renumbers(tmp_path):
    f = build_flat(tmp_path)
    plan = tmp_path / "plan.yaml"
    plan.write_text(PLAN, encoding="utf-8")
    r = run("h3group", f, "--plan", plan, "--apply")
    assert r.returncode == 0, r.stdout + r.stderr
    doc = B.Doc(f)
    heads = [(doc.level(k), B.ptext(k).strip()) for k in doc.body if doc.is_heading(k)]
    assert heads[4:] == [(2, "7.2 实施方案"), (3, "7.2.1 总体思路"), (3, "7.2.2 布局与工程")]
    assert items_under(doc) == ["1）方案框架", "2）规划约束", "1）供需平衡", "2）空间布局",
                                "3）重点工程", "4）保障措施"]
    s = seq(f)
    assert s[s.index("TBL") - 1] == "表7-1 指" and s[-1] == "SECT"
    assert s[s.index("4）保障措施") + 1].startswith("新增正文")
    r = run("check", f).stdout
    assert "标题样式N）: 0" in r and "条目编号: 0" in r and "标题跳级: 0" in r, r
    assert run("h3group", f, "--plan", plan, "--apply").returncode != 0      # 重复执行必须拒绝


def test_h3group_renumbers_legacy_heading_items(tmp_path):
    d = Document()
    d.add_heading("7 技术路线", level=1)
    d.add_heading("7.1 现状调查", level=2)
    d.add_heading("7.1.1 思路", level=3)
    d.add_paragraph("样板正文" * 20)
    d.add_heading("7.2 实施方案", level=2)
    for i, name in enumerate(["方案框架", "规划约束", "供需平衡"], 1):
        d.add_heading(f"{i}）{name}", level=4)          # 旧稿：N）用了四级标题样式
        d.add_paragraph(f"{name}正文" * 15)
    f = tmp_path / "legacy.docx"
    d.save(f)
    plan = tmp_path / "plan.yaml"
    plan.write_text("""
sections:
  - h2: "7.2 实施方案"
    groups:
      - {title: "7.2.1 总体思路", n: 1}
      - {title: "7.2.2 布局", n: 2, append: [{item: "保障措施", paras: ["新增正文。"]}]}
""", encoding="utf-8")
    r = run("h3group", f, "--plan", plan, "--apply")
    assert r.returncode == 0, r.stdout + r.stderr
    doc = B.Doc(f)
    heads = [B.ptext(k).strip() for k in doc.body if doc.is_heading(k) and doc.level(k) == 4]
    assert heads == ["1）方案框架", "1）规划约束", "2）供需平衡"]
    assert items_under(doc) == ["3）保障措施"]                # 追加条目是正文段，不再克隆四级标题
    r = run("check", f)
    assert r.returncode == 2 and "标题样式N）: 3" in r.stdout   # 旧稿标题样式仍报红，待按体例改


def test_check_flags_heading_items_and_body_gaps(tmp_path):
    d = Document()
    d.add_heading("6.1 概况", level=2)
    d.add_heading("6.1.1 依据", level=3)
    d.add_paragraph("法规类依据如下。")
    d.add_paragraph("1）《水法》")
    d.add_paragraph("2）《水污染防治法》")
    d.add_paragraph("规划类依据如下。")
    d.add_paragraph("1）《海宁水网建设规划》")              # 同一标题下第二组清单从 1）重起，不算断点
    d.add_paragraph("2）《海宁市域污水工程专项规划》")
    f = tmp_path / "ok.docx"
    d.save(f)
    r = run("check", f)
    assert r.returncode == 0, r.stdout
    d.add_paragraph("4）跳号条目")
    d.add_heading("6.1.1.1 实施步骤", level=4)
    d.add_paragraph("1）第一步")                              # 标题后清零，从 1）起合规
    d.add_heading("（十九）旧式条目", level=4)
    d.add_heading("20）旧式条目", level=4)
    d.save(f)
    r = run("check", f)
    assert r.returncode == 2
    assert "标题样式N）: 2" in r.stdout and "条目编号: 1" in r.stdout and "2→4" in r.stdout


def test_h4num_is_retired(tmp_path):
    f = build(tmp_path)
    before = f.read_bytes()
    r = run("h4num", f, "--apply")
    assert r.returncode != 0 and "已退役" in r.stderr
    assert f.read_bytes() == before


def test_h3group_refuses_count_mismatch(tmp_path):
    f = build_flat(tmp_path)
    plan = tmp_path / "plan.yaml"
    plan.write_text(PLAN.replace("n: 3", "n: 2"), encoding="utf-8")
    before = f.read_bytes()
    assert run("h3group", f, "--plan", plan, "--apply").returncode != 0
    assert f.read_bytes() == before


def build_cited(tmp_path):
    d = Document()
    d.add_heading("6.1 概况", level=2)
    d.add_paragraph("水网规划记载总设计规模7万m³/日（《海宁水网建设规划》第10页）。本项目拟采用分时段方法，原拟用途保留。")
    d.add_paragraph("规划期为（2021—2035年），以《管理办法》为指引（招标文件第二章“招标需求”）；并开展模拟。")
    d.add_paragraph("表中事实定位分别见（《海宁水网建设规划》第10页；《海宁市域污水工程专项规划修编（2022—2035）》第150页）。矩阵用于说明资料如何进入分析。")
    out = tmp_path / "c.docx"
    d.save(out)
    return out


def test_uncite_removes_source_parens_only(tmp_path):
    f = build_cited(tmp_path)
    assert "正文出处括注: 3" in run("check", f).stdout
    assert run("uncite", f, "--apply").returncode == 0
    ts = [B.ptext(p) for p in B.Doc(f).root.iter(B.w("p"))]
    assert ts[1].startswith("水网规划记载总设计规模7万m³/日。")
    assert "（2021—2035年）" in ts[2] and "以《管理办法》为指引；" in ts[2]      # 非出处括注不动
    assert ts[3] == "矩阵用于说明资料如何进入分析。"                           # “分别见（…）。”整句删
    assert "正文出处括注: 0" in run("check", f).stdout


def test_resub_keeps_protected_terms_and_numbers(tmp_path):
    f = build_cited(tmp_path)
    rules = tmp_path / "r.yaml"
    rules.write_text('resub:\n  - ["(?<![模原])拟(?!议|定)", ""]\n', encoding="utf-8")
    assert run("resub", f, "--rules", rules, "--apply").returncode == 0
    t = "".join(B.ptext(p) for p in B.Doc(f).root.iter(B.w("p")))
    assert "本项目采用分时段方法" in t and "原拟用途" in t and "模拟" in t
    rules.write_text('resub:\n  - ["7万", "8万"]\n', encoding="utf-8")
    before = f.read_bytes()
    assert run("resub", f, "--rules", rules, "--apply").returncode != 0           # 改数字必须被护栏拦下
    assert f.read_bytes() == before


def test_media_swaps_by_hash_and_rejects_size_change(tmp_path):
    from PIL import Image
    f = build(tmp_path)
    old, new = tmp_path / "old", tmp_path / "new"
    old.mkdir(); new.mkdir()
    (old / "x.png").write_bytes(PNG)
    Image.new("RGB", (1, 1), "red").save(new / "x.png")
    r = run("media", f, "--old", old, "--new", new, "--apply")
    assert r.returncode == 0 and "替换 1 张" in r.stdout
    import zipfile
    z = zipfile.ZipFile(f)
    assert any(z.read(n) == (new / "x.png").read_bytes() for n in z.namelist() if n.startswith("word/media/"))
    f2 = build(tmp_path)
    Image.new("RGB", (2, 2), "red").save(new / "x.png")
    assert run("media", f2, "--old", old, "--new", new, "--apply").returncode != 0


def test_swapfig_changes_ratio_caption_and_drops_old_part(tmp_path):
    """海宁第7章重画 10 张图：新图比例与旧图不同、图题要改名；旧图部件不能留成孤儿（health gate 判红）。"""
    import json
    import zipfile
    from PIL import Image
    f = build(tmp_path)
    new = tmp_path / "f7.png"
    Image.new("RGB", (400, 200), "blue").save(new)
    plan = tmp_path / "plan.json"
    plan.write_text(json.dumps([{"no": "1-1", "png": str(new), "caption": "调查成果复核流程图"}], ensure_ascii=False))
    old_media = [n for n in zipfile.ZipFile(f).namelist() if n.startswith("word/media/")]
    r = run("swapfig", f, "--plan", plan, "--apply")
    assert r.returncode == 0, r.stderr
    z = zipfile.ZipFile(f)
    names = z.namelist()
    assert not set(old_media) & set(names)
    assert any(z.read(n) == new.read_bytes() for n in names if n.startswith("word/media/"))
    xml = z.read("word/document.xml").decode()
    assert 'cx="5580000" cy="2790000"' in xml               # 15.5cm 宽、按 2:1 比例
    assert "图1-1 调查成果复核流程图" in [B.ptext(p).strip() for p in B.Doc(f).body.iter(B.w("p"))]
    assert seq(f).count("IMG") == 1 and seq(f)[-1] == "SECT"


def test_swapfig_refuses_unknown_figure(tmp_path):
    import json
    f = build(tmp_path)
    plan = tmp_path / "plan.json"
    plan.write_text(json.dumps([{"no": "9-9", "png": str(tmp_path / "x.png")}]))
    assert run("swapfig", f, "--plan", plan, "--apply").returncode != 0



# ── insert：定稿扩充按锚点原文插段（2026-09-18 海宁 6.2.3）────────────────
def test_insert_keeps_originals_and_is_idempotent(tmp_path):
    src = build(tmp_path)
    patch = tmp_path / "p.md"
    patch.write_text("@after 第一段正文\n补充甲。\n补充乙。\n@after 1.1.2 资源约束\n补充丙。\n", encoding="utf-8")
    before = [B.ptext(p) for p in B.Doc(src).body.iter(B.w("p"))]
    run = lambda: subprocess.run([sys.executable, str(TOOL), "insert", str(src), "--patch", str(patch), "--apply"],
                                 capture_output=True, text=True)
    r = run()
    assert r.returncode == 0, r.stderr
    s = seq(src)
    assert s[s.index("第一段正文，")+1:s.index("第一段正文，")+3] == ["补充甲。", "补充乙。"]
    assert s[s.index("1.1.2 ")+1] == "补充丙。"          # 锚点是标题：插在标题后
    after = [B.ptext(p) for p in B.Doc(src).body.iter(B.w("p"))]
    it = iter(after)
    assert all(any(t == n for n in it) for t in before)        # 原段按序全在
    assert "TBL" in s and s[-1] == "SECT"
    assert "共插入 0 段" in run().stdout                         # 重跑不重复插


def test_insert_rejects_ambiguous_anchor(tmp_path):
    src = build(tmp_path)
    patch = tmp_path / "p.md"
    patch.write_text("@after 第\n补充。\n", encoding="utf-8")
    r = subprocess.run([sys.executable, str(TOOL), "insert", str(src), "--patch", str(patch), "--apply"],
                       capture_output=True, text=True)
    assert r.returncode != 0 and "命中" in r.stderr
