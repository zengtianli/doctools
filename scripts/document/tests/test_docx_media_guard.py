"""图片守卫回归门 —— 「重建文档把内嵌图全丢了还覆盖了原稿」必须被拦在 replace 之前。

守卫在 `lib/docx_safe_save.py::_check_media` 与 `lib/docx_parts.py::assert_media_intact`。
每条「拦住」用例都同时断言**原文件字节未变**：先覆盖再报错等于没有安全网。

回退即红：把 `lib/docx_safe_save.py` 换回加守卫前的版本，`test_rebuild_overwrite_*` /
`test_inplace_delete_*` / `test_cross_document_*` / `test_guard_off_is_red_proof` 四条变红（2026-09-17 实测）。
`test_guard_off_is_red_proof` 另在子进程里用 `DOCX_GRAFT_OFF=1` 常驻对照：关掉收口，原稿真的被覆盖成 0 图。
"""
from __future__ import annotations

import hashlib
import os
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "lib"))
sys.path.insert(0, str(ROOT / "tests" / "smoke"))

import docx_safe_save  # noqa: E402
from docx_parts import (PartIntegrityError, assert_parts_intact,  # noqa: E402
                        media_census)
from _fixture import make_png  # noqa: E402  stdlib PNG 造件器

pytestmark = pytest.mark.skipif(
    not docx_safe_save.patched, reason="DOCX_GRAFT_OFF=1：收口没装，守卫不在岗")

N_IMG = 3


def _md5(p: Path) -> str:
    return hashlib.md5(p.read_bytes()).hexdigest()


def _make(tmp_path: Path, n: int = N_IMG, name: str = "src.docx") -> Path:
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    doc.add_paragraph("第一段正文")
    for i in range(n):
        png = make_png(tmp_path / f"img{i}.png", rgb=(40 * i + 10, 90, 200 - 30 * i))
        doc.add_picture(str(png), width=Cm(3))
        doc.add_paragraph(f"图 {i + 1} 之后的文字")
    out = tmp_path / name
    doc.save(str(out))
    return out


def _rebuild_text_only(src: Path):
    """事故原型：读旧稿 → Document() 新建 → 只搬文字。"""
    from docx import Document

    old, new = Document(str(src)), Document()
    for p in old.paragraphs:
        new.add_paragraph(p.text)
    return new


def test_fixture_really_has_images(tmp_path):
    c = media_census(_make(tmp_path))
    assert c == {"media": N_IMG, "refs": N_IMG, "dangling": []}


def test_normal_edit_passes(tmp_path):
    from docx import Document

    src = _make(tmp_path)
    doc = Document(str(src))
    doc.paragraphs[0].text = "改过的第一段"
    doc.save(str(src))
    assert media_census(src)["refs"] == N_IMG
    assert Document(str(src)).paragraphs[0].text == "改过的第一段"


def test_rebuild_overwrite_is_blocked_and_original_untouched(tmp_path):
    src = _make(tmp_path)
    before = _md5(src)
    new = _rebuild_text_only(src)
    with pytest.raises(PartIntegrityError, match="图片守卫"):
        new.save(str(src))
    assert _md5(src) == before, "守卫抛错了但原文件已经被覆盖"
    assert not list(tmp_path.glob("*.dsafe")), "临时件没收掉"


def test_rebuild_to_new_path_is_not_blocked(tmp_path):
    """新路径没有东西被覆盖 —— 从零造文件（pdf_to_docx 这类）不许被误拦。"""
    src = _make(tmp_path)
    out = tmp_path / "new.docx"
    _rebuild_text_only(src).save(str(out))
    assert media_census(out)["refs"] == 0


def test_inplace_delete_picture_blocked_then_allowed(tmp_path):
    from docx import Document

    src = _make(tmp_path)
    before = _md5(src)

    def drop_first_picture():
        doc = Document(str(src))
        for p in doc.paragraphs:
            if p._p.xpath(".//w:drawing"):
                p._p.getparent().remove(p._p)
                break
        return doc

    with pytest.raises(PartIntegrityError, match="正文图引用 3 → 2"):
        drop_first_picture().save(str(src))
    assert _md5(src) == before

    with docx_safe_save.allow_media_loss():
        drop_first_picture().save(str(src))
    assert media_census(src)["refs"] == N_IMG - 1


def test_cross_document_copy_dangling_blocked(tmp_path):
    """把含图的 body 元素 deepcopy 进新文档：r:embed 指向新包里不存在的关系。"""
    import copy

    from docx import Document

    src = _make(tmp_path)
    old, new = Document(str(src)), Document()
    body = new.element.body
    for el in old.element.body.iterchildren():
        if el.tag.endswith("}p"):
            body.insert(len(body) - 1, copy.deepcopy(el))
    out = tmp_path / "copied.docx"
    with pytest.raises(PartIntegrityError, match="悬空图片引用"):
        with docx_safe_save.allow_media_loss():      # 悬空不受 opt-out 放行
            new.save(str(out))
    assert not out.exists()


def test_zero_image_doc_is_clean(tmp_path):
    from docx import Document

    src = _make(tmp_path, n=0)
    assert media_census(src) == {"media": 0, "refs": 0, "dangling": []}
    doc = Document(str(src))
    doc.add_paragraph("追加")
    doc.save(str(src))


def test_census_refuses_non_docx(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(PartIntegrityError):
        media_census(bad)


def _strip_one_drawing(src: Path, dst: Path, *, break_rel: bool = False) -> None:
    import re

    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "word/document.xml" and not break_rel:
                data, n = re.subn(rb"<w:drawing>.*?</w:drawing>", b"", data, count=1, flags=re.S)
                assert n == 1
            if item.filename == "word/document.xml" and break_rel:
                data, n = re.subn(rb'r:embed="[^"]+"', b'r:embed="rId9999"', data, count=1)
                assert n == 1
            zo.writestr(item, data)


def test_surgical_ref_loss_modes(tmp_path, capsys):
    src = _make(tmp_path)
    dst = tmp_path / "out.docx"
    _strip_one_drawing(src, dst)

    assert_parts_intact(src, dst, verbose=False)               # 默认：告警不抛
    assert "正文图引用 3 → 2" in capsys.readouterr().err
    with pytest.raises(PartIntegrityError, match="正文图引用 3 → 2"):
        assert_parts_intact(src, dst, verbose=False, allow_media_loss=False)
    assert_parts_intact(src, dst, verbose=False, allow_media_loss=True)
    assert "正文图引用" not in capsys.readouterr().err


def test_surgical_dangling_always_raises(tmp_path):
    src = _make(tmp_path)
    dst = tmp_path / "out.docx"
    _strip_one_drawing(src, dst, break_rel=True)
    with pytest.raises(PartIntegrityError, match="悬空图片引用"):
        assert_parts_intact(src, dst, verbose=False, allow_media_loss=True)


def test_guard_off_is_red_proof(tmp_path):
    """关掉收口（DOCX_GRAFT_OFF=1）后同一个重建动作会把原稿覆盖成 0 图 ——
    证明拦住它的是守卫本身，不是别的巧合。"""
    src = _make(tmp_path)
    code = (
        "import sys; sys.path.insert(0, sys.argv[1])\n"
        "import docx_safe_save\n"
        "from docx import Document\n"
        "old, new = Document(sys.argv[2]), Document()\n"
        "[new.add_paragraph(p.text) for p in old.paragraphs]\n"
        "new.save(sys.argv[2])\n"
    )
    env = {**os.environ, "DOCX_GRAFT_OFF": "1", "DOCX_GRAFT_QUIET": "1"}
    r = subprocess.run([sys.executable, "-c", code, str(ROOT / "lib"), str(src)],
                       env=env, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    assert media_census(src)["refs"] == 0          # 图真的全丢了

    src2 = _make(tmp_path, name="src2.docx")
    env.pop("DOCX_GRAFT_OFF")
    r = subprocess.run([sys.executable, "-c", code, str(ROOT / "lib"), str(src2)],
                       env=env, capture_output=True, text=True)
    assert r.returncode != 0 and "图片守卫" in r.stderr
    assert media_census(src2)["refs"] == N_IMG
