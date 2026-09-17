#!/usr/bin/env python3
"""docx_safe_save — import 它一次，本进程里所有 python-docx 存盘自动收口成 surgical。

    import docx_safe_save  # noqa: F401

## 它解决什么（2026-07-30 立，数据是实测的）

一份 301 部件 / 75 个嵌入公式的真报告，python-docx **什么都不改**地打开再存回：

    zip 条目数     301 → 301        一个没少
    字节变了的部件  60 个            含 [Content_Types].xml / _rels/.rels / 28 header / 14 footer
    语义变了的部件  1 个             只有 [Content_Types].xml 是真的变了

也就是 59/60 纯属被 python-docx 的序列化器重写了一遍，内容一模一样。
**光数条目数会得出「python-docx 没损失」的结论 —— 那是拿错尺子。** 真该量的是
「这次改一个字号，到底碰了多少个不该碰的文件」：60 个，每个都是一次
「Word 可能渲染得不一样」的机会（属性序 / 命名空间声明位置 / 空白差异）。

surgical（lxml + zipfile 只重写点名的部件）炸开面是 **1**。

## 为什么是补 `OpcPackage`，不是改 36 个脚本

doctools 里 36 个非测试脚本用 python-docx 存盘、合计 8500+ 行。改写成手写 XML =
重造 36 个轮子。但实测（`inspect.getsource`）它们的存盘路径**只有一条**：

    Document.save → DocumentPart.save → OpcPackage.save → PackageWriter.write
    Document(path) → Package.open  →   OpcPackage.open

两个都是**类级别**的方法，所以补在这里：① 一处覆盖全部调用方 ② 不受 import 顺序
影响（脚本顶部先 `from docx import Document` 也照样生效，因为实例共享类）
③ docxcompose 之类第三方封装同样落进这条路。

## 补完之后 save 干了什么

    1. python-docx 原样写到同目录临时文件（不碰目标）
    2. graft_unchanged(源文件, 临时文件)：逐部件做 XML 规范化(C14N)对比，
       **语义没变的按源文件原始字节还原**
    3. 原子 replace 到目标

顺序是刻意的：graft 在 replace **之前**。所以万一 graft 判定这次改动把部件整个弄丢了
（那不是「改」是「重造」），抛错时目标文件还是原样 —— 先写坏再报错等于没有安全网。

新建文档（`Document()` 无参，源就是 python-docx 自带模板）**不 graft** —— 没有
「原件」可保留，pdf_to_docx 这类从零造文件的走这条。

## 图片守卫（2026-09-17 加）

部件集合比对看不见「重建文档把内嵌图全丢了」：`Document()` 新建再把文字搬过去、
覆盖存回原路径，走的正是上面那条不 graft 的分支。所以不论有没有源，临时件写好之后、
replace 之前再量一次图（`docx_parts.media_census`）：

    目标路径上已有一份 docx → 以**即将被覆盖的那份**为基线，图片部件数或正文图引用数
                              变少 = 抛 PartIntegrityError，目标文件一个字节不动
    新出现的悬空图片关系     → 抛（跨文档搬元素的症状）
    存到新路径且比源件图少   → 只打印告警（拆册/取一章是正常操作，源件没被覆盖）

范围：只管 import 了本模块的进程。临时脚本裸用 python-docx 不经过这里。

## 逃生 / 调参

    DOCX_GRAFT_OFF=1     完全不打补丁（退回裸 python-docx 行为）
    DOCX_GRAFT_QUIET=1   不打印每次收口的那行 stderr
    with docx_safe_save.allow_part_loss(): ...   这段里允许部件丢失（默认抛错）
    with docx_safe_save.allow_media_loss(): ...  这段里允许覆盖存盘后图变少（默认抛错）
"""
from __future__ import annotations

import contextlib
import os
import sys
from pathlib import Path
from weakref import WeakKeyDictionary

# append 不 insert(0)：lib/ 和 sub/ 有同名模块(styles)，插在 0 位会顶掉
# 脚本自己那一份。append 让「顶掉」在结构上不可能发生。
sys.path.append(str(Path(__file__).resolve().parent))
from docx_surgical import graft_unchanged  # noqa: E402
from docx_parts import (PartIntegrityError, assert_media_intact,  # noqa: E402
                        media_census)

__all__ = ["allow_part_loss", "allow_media_loss", "patched"]

_QUIET = os.environ.get("DOCX_GRAFT_QUIET") == "1"
_SENTINEL = "_docx_safe_save_patched"

# package 对象 → 它是从哪个文件读进来的。WeakKey：文档被回收后自动清，不留内存。
_SRC: WeakKeyDictionary = WeakKeyDictionary()

# 允许部件丢失的作用域深度（嵌套安全）
_ALLOW_LOSS = 0


@contextlib.contextmanager
def allow_part_loss():
    """这段代码里的存盘允许丢部件（默认丢部件 = 抛 RepackError）。

    只在**确实是「重造一份新文件」而不是「改一份现有文件」**时用 —— 丢部件的典型
    后果是：正文里那个公式对象的关系还在，指向的部件却没了，Word 打开报「内容有问题」。
    """
    global _ALLOW_LOSS
    _ALLOW_LOSS += 1
    try:
        yield
    finally:
        _ALLOW_LOSS -= 1


# 允许覆盖存盘后图变少的作用域深度（嵌套安全）
_ALLOW_MEDIA_LOSS = 0


@contextlib.contextmanager
def allow_media_loss():
    """这段代码里的存盘允许「覆盖一份 docx 后图比原来少」（默认 = 抛 PartIntegrityError）。

    只在**有意删图**时用（删含图的块、原地瘦身）。新出现的悬空图片关系不受它放行。
    """
    global _ALLOW_MEDIA_LOSS
    _ALLOW_MEDIA_LOSS += 1
    try:
        yield
    finally:
        _ALLOW_MEDIA_LOSS -= 1


def _census_or_none(path: Path):
    """基线普查。目标路径上的旧文件不是 docx（空文件/别的格式）→ 没有图可丢，返回 None。"""
    try:
        return media_census(path)
    except PartIntegrityError:
        return None


def _check_media(src, target: Path, tmp: Path) -> None:
    """replace 之前调。抛错 = 目标文件未动。"""
    after = media_census(tmp)
    old = _census_or_none(target) if target.is_file() else None
    origin = _census_or_none(src) if src is not None else None
    # 悬空引用的基线：源件本来就有的不算这次弄坏的
    base_dangling = (origin or old or {"dangling": []})["dangling"]
    if old is not None:
        assert_media_intact({**old, "dangling": base_dangling}, after,
                            allow_loss=bool(_ALLOW_MEDIA_LOSS),
                            label=f"覆盖 {target.name}")
    else:
        assert_media_intact({"media": 0, "refs": 0, "dangling": base_dangling}, after,
                            label=f"写出 {target.name}")
        if (origin is not None and after["refs"] < origin["refs"]
                and not _ALLOW_MEDIA_LOSS and not _QUIET):
            print(f"⚠ [图片守卫] {target.name}: 正文图引用 {origin['refs']} → "
                  f"{after['refs']}（比源件 {Path(src).name} 少；存的是新路径，未拦）",
                  file=sys.stderr)


def _default_template() -> Path | None:
    """python-docx 自带的空模板路径。从它打开的文档 = 新建，没有「原件」要保。"""
    try:
        from docx.api import _default_docx_path
        return Path(_default_docx_path()).resolve()
    except Exception:
        return None


def _as_path(x) -> Path | None:
    """只认真实的文件系统路径。file-like / BytesIO 一律返回 None（不 graft）。"""
    if isinstance(x, (str, os.PathLike)):
        try:
            return Path(x)
        except (TypeError, ValueError):
            return None
    return None


def _install() -> bool:
    if os.environ.get("DOCX_GRAFT_OFF") == "1":
        if not _QUIET:
            print("⚠ DOCX_GRAFT_OFF=1 → 不打 surgical 收口补丁，python-docx 裸存盘"
                  "（炸开面 ~60 个部件）", file=sys.stderr)
        return False

    from docx.opc.package import OpcPackage

    if getattr(OpcPackage, _SENTINEL, False):
        return True  # 已补过（同进程重复 import 走这里）

    tmpl = _default_template()
    _orig_open = OpcPackage.open.__func__      # 拆 classmethod 拿到底层函数
    _orig_save = OpcPackage.save

    def _open(cls, pkg_file):
        pkg = _orig_open(cls, pkg_file)
        p = _as_path(pkg_file)
        if p is not None:
            try:
                rp = p.resolve()
            except OSError:
                return pkg
            if rp.is_file() and rp != tmpl:
                _SRC[pkg] = rp
        return pkg

    def _save(self, pkg_file):
        src = _SRC.get(self)
        target = _as_path(pkg_file)
        if target is None:
            # 存到流里 —— 没有落盘路径，没有东西会被覆盖，原样存。
            return _orig_save(self, pkg_file)

        if target.is_symlink():
            # os.replace 换掉的是链接本身；写到链接指向的真文件上，和裸 python-docx 行为一致。
            target = target.resolve()
        tmp = target.with_name(target.name + ".dsafe")
        try:
            _orig_save(self, str(tmp))
            # 新建文档没有原件可 graft，但照样要过图片守卫：
            # 「Document() 重建后覆盖原稿」正是这条分支。
            info = ({"还原": [], "真变了": [], "新增": []} if src is None else
                    graft_unchanged(src, tmp,
                                    on_missing="ignore" if _ALLOW_LOSS else "error"))
            _check_media(src, target, tmp)
            os.replace(tmp, target)
        except BaseException:
            # 目标文件至今一个字节都没动过；把临时件收掉，让错误照原样往上抛。
            with contextlib.suppress(OSError):
                tmp.unlink()
            raise

        if not _QUIET and info["还原"]:
            real = info["真变了"]
            shown = "、".join(real[:3]) + (f" 等 {len(real)} 个" if len(real) > 3 else "")
            print(f"[surgical] {target.name}: 只重写 {len(real)} 个部件"
                  f"（{shown or '无'}），还原 {len(info['还原'])} 个被重新序列化的部件"
                  + (f"，新增 {len(info['新增'])} 个" if info["新增"] else ""),
                  file=sys.stderr)
        return None

    OpcPackage.open = classmethod(_open)
    OpcPackage.save = _save
    setattr(OpcPackage, _SENTINEL, True)
    return True


patched = _install()
