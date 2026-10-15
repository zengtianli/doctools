from pathlib import Path
import sys
import zipfile

from docx import Document
from docx.shared import Pt
from lxml import etree
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'lib'))
import docx_reference as R


def make(path, body_size):
    d=Document();d.add_paragraph('示范标题').runs[0].font.size=Pt(22)
    d.add_heading('一、主要任务',1)
    d.add_paragraph('正文保留。').runs[0].font.size=Pt(body_size)
    d.add_paragraph('1. 建设1处站点。');d.add_paragraph('2. 完成3次比测。')
    t=d.add_table(rows=2,cols=2);t.cell(0,0).text='项目';t.cell(0,1).text='水量'
    t.cell(1,0).text='目标';t.cell(1,1).text='0.14 m³/s'
    d.save(path)


def test_surgical_format_keeps_tables_package_and_input(tmp_path):
    source=tmp_path/'src.docx';ref=tmp_path/'ref.docx';out=tmp_path/'out.docx'
    make(source,12);make(ref,16);before=source.read_bytes();reference=ref.read_bytes()
    result=R.apply(source,ref,out,assignments=[{'block':0,'role':'title'},{'block':2,'role':'body'}],
                   roles={'title':0,'body':2},tables=False)
    assert source.read_bytes()==before and ref.read_bytes()==reference
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(out) as b:
        assert a.namelist()==b.namelist()
        assert all(a.read(n)==b.read(n) for n in a.namelist() if n!='word/document.xml')
        root=etree.fromstring(b.read('word/document.xml'))
        assert ''.join(root.itertext()).find('0.14 m³/s')>=0
    assert result['block_mapping']['5']==[5]
    with pytest.raises(ValueError,match='新文件'): R.apply(source,ref,out,assignments=[],roles={'body':2})


def test_merge_remaps_all_following_blocks_and_refuses_tables(tmp_path):
    src=tmp_path/'s.docx';ref=tmp_path/'r.docx';make(src,12);make(ref,16)
    result=R.apply(src,ref,tmp_path/'merged.docx',assignments=[],roles={'body':2},tables=False,
        replacements=[{'start':3,'end':4,'paragraphs':[{'text':'一是建设1处站点。二是完成3次比测。'}]}])
    assert result['block_mapping']['3']==[3] and result['block_mapping']['4']==[3]
    assert result['block_mapping']['5']==[4]
    with pytest.raises(ValueError,match='表格'):
        R.apply(src,ref,tmp_path/'bad.docx',assignments=[],roles={'body':2},
                replacements=[{'start':4,'end':5,'paragraphs':[{'text':'不能吞表格'}]}])
    assert not (tmp_path/'bad.docx').exists()


def test_partial_latin_override_keeps_inherited_chinese_font(tmp_path):
    path=tmp_path/'ref.docx';d=Document()
    normal=d.styles['Normal'];normal.font.name='仿宋_GB2312'
    normal.element.rPr.rFonts.set(R.qn('w:eastAsia'),'仿宋_GB2312')
    d.add_paragraph('标题')
    run=d.add_paragraph('中文正文').runs[0];run.font.name='Times New Roman';run.font.size=Pt(16)
    d.save(path)
    shell=R.shell(path,1)
    assert shell['r']['rFonts'].get(R.qn('w:eastAsia'))=='仿宋_GB2312'
    assert shell['r']['rFonts'].get(R.qn('w:ascii'))=='Times New Roman'
