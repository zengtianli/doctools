#!/usr/bin/env python3
"""Exercise real GUI backend failure envelopes and subsequent successful work."""
from pathlib import Path

from docx import Document

from _common import backend, finish, sha256, workspace


def expect_error(*args: str) -> dict:
    result = backend(*args)
    assert result["ok"] is False, result
    assert isinstance(result.get("error"), str) and result["error"].strip(), result
    return result


def assert_success(result: dict, source: Path, root: Path) -> None:
    assert result["ok"] is True and result["succeeded"] == 1, result
    rows = [row for row in result["results"] if row["input"] == str(source)]
    assert len(rows) == 1 and rows[0]["ok"] is True, result
    outputs = [Path(path) for path in rows[0]["outputs"]]
    assert outputs, result
    for output in outputs:
        assert output.is_relative_to(root) and output.is_file(), output
        assert output != source, output
    texts = ["".join(p.text for p in Document(output).paragraphs)
             for output in outputs if output.suffix == ".docx"]
    assert any("“恢复”" in text and "20m²" in text for text in texts), texts


def main() -> None:
    cases = []
    with workspace("recovery") as raw:
        root = Path(raw)
        good = root / "正常输入.docx"
        retry = root / "失败后重试.docx"
        bad = root / "损坏输入.docx"
        missing = root / "不存在.docx"
        for path in (good, retry):
            doc = Document()
            doc.add_paragraph('测试, "恢复" 20平方米')
            doc.save(path)
        bad.write_bytes(b"invalid DOCX acceptance fixture\n")
        originals = {path: sha256(path) for path in (good, retry, bad)}

        # Read option IDs from the same runtime catalog used by the app.
        catalog = backend("gui-ops")
        clean = next(op for op in catalog["ops"] if op["id"] == "clean")
        bool_option = next(option["id"] for option in clean["options"]
                           if option["type"] == "bool")
        invalid_calls = [
            ("unknown_operation", ["--op", "acceptance-unknown"]),
            ("unknown_option", ["--op", "clean", "--opt", "acceptance-unknown=1"]),
            ("invalid_bool", ["--op", "clean", "--opt", f"{bool_option}=invalid"]),
            ("malformed_option", ["--op", "clean", "--opt", "no-equals-sign"]),
            ("missing_conversion_target", ["--op", "convert"]),
            ("missing_required_template", ["--op", "formatclone"]),
        ]
        for name, args in invalid_calls:
            expect_error("gui-run", *args, "--files", good)
            assert {path: sha256(path) for path in originals} == originals, name
            assert set(root.iterdir()) == set(originals), name
            cases.append(name)

        expect_error("gui-run", "--op", "clean", "--files", missing)
        assert not missing.exists()
        cases.append("missing_file")

        corrupt = backend("gui-run", "--op", "clean", "--files", bad)
        assert corrupt["ok"] is False, f"Corrupt DOCX falsely reported success: {corrupt}"
        assert corrupt["succeeded"] == 0 and corrupt["total"] == 1, corrupt
        assert corrupt["results"][0]["ok"] is False, corrupt
        assert corrupt["results"][0]["outputs"] == [], corrupt
        assert "处理失败" in corrupt["results"][0]["message"], corrupt
        assert corrupt["log"].strip(), corrupt
        cases.append("corrupt_file_failure")

        mixed = backend("gui-run", "--op", "clean", "--files", bad, good, missing)
        assert mixed["total"] == 2 and len(mixed["results"]) == 2, mixed
        assert mixed["skipped_missing"] == [missing.name], mixed
        failed = next(row for row in mixed["results"] if row["input"] == str(bad))
        assert failed["ok"] is False and failed["outputs"] == [], mixed
        assert_success(mixed, good, root)
        cases.append("mixed_batch_continues_after_failure")

        recovered = backend("gui-run", "--op", "clean", "--files", retry)
        assert recovered["total"] == 1, recovered
        assert_success(recovered, retry, root)
        assert {path: sha256(path) for path in originals} == originals
        cases.extend(["next_request_recovers", "all_source_hashes_preserved"])

    finish("recovery", "真实 gui-run 错误信封、批次隔离及失败后恢复通过",
           cases=cases, case_count=len(cases),
           scope="临时 DOCX 输入；不涵盖进程强杀、系统断电或磁盘故障恢复")


if __name__ == "__main__":
    main()
