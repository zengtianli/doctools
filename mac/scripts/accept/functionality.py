#!/usr/bin/env python3
"""Exercise DocKit's actual GUI JSON entry points with isolated document fixtures."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from _common import backend, cli, finish, sha256, workspace


def successful(payload: dict, source: Path, root: Path) -> list[Path]:
    assert payload["ok"] is True, payload
    assert payload["succeeded"] == payload["total"] == 1, payload
    result = payload["results"][0]
    assert result["ok"] and result["message"], payload
    assert payload["log"].strip(), "The GUI must receive the processing report"
    outputs = [Path(item) for item in result["outputs"]]
    assert outputs, payload
    for output in outputs:
        assert output.exists() and output.resolve().is_relative_to(root.resolve()), output
        assert output.resolve() != source.resolve(), "Copy-producing operation overwrote source"
    return outputs


def main() -> None:
    catalog = backend("gui-ops")
    assert catalog["ok"] and catalog["ops"], catalog
    ops = {op["id"]: op for op in catalog["ops"]}
    assert len(ops) == len(catalog["ops"]), "Operation IDs must be unique"
    assert {"clean", "quotes", "merge", "convert", "formatclone"} <= ops.keys()
    option_count = 0
    for op in ops.values():
        assert all(isinstance(op[key], str) and op[key] for key in ("id", "title", "subtitle", "icon"))
        assert op["kind"] in ("files", "dir")
        assert isinstance(op["exts"], list)
        assert op["kind"] == "dir" or op["exts"]
        options = op.get("options", [])
        groups = {group["id"] for group in op.get("option_groups", [])}
        assert len({option["id"] for option in options}) == len(options)
        for option in options:
            assert option["group"] in groups and option["title"]
            assert option["type"] in ("bool", "file")
            if option["type"] == "bool":
                assert isinstance(option["default"], bool)
        option_count += len(options)

    checked: list[str] = ["gui-ops catalog and declarative option contract"]
    with workspace("functionality") as temp:
        root = Path(temp)
        source = root / "中文 空格验收.docx"
        body = '正文:"成章",面积10平方米,小时候。'
        table = '表格:"保留",面积20平方米。'
        header = '页眉:"标题",面积30平方米。'
        doc = Document()
        doc.add_paragraph(body)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = table
        doc.sections[0].header.paragraphs[0].text = header
        doc.sections[0].footer.paragraphs[0].text = "固定页脚"
        doc.save(source)
        original = sha256(source)

        payload = backend("gui-run", "--op", "clean", "--files", str(source))
        outputs = successful(payload, source, root)
        output = next(path for path in outputs if path.suffix == ".docx")
        cleaned = Document(output)
        assert cleaned.paragraphs[0].text == "正文：“成章”，面积10m²，小时候。", cleaned.paragraphs[0].text
        assert cleaned.tables[0].cell(0, 0).text == "表格：“保留”，面积20m²。"
        assert cleaned.sections[0].header.paragraphs[0].text == "页眉：“标题”，面积30m²。"
        assert cleaned.sections[0].footer.paragraphs[0].text == "固定页脚"
        assert cleaned.sections[0]._sectPr.find(qn("w:headerReference")) is not None
        assert cleaned.sections[0]._sectPr.find(qn("w:footerReference")) is not None
        assert sha256(source) == original
        checked.append("DOCX clean: body/table/header normalization, unit boundaries, header/footer retention and byte-identical source")

        selective = root / "仅正文.docx"
        selective.write_bytes(source.read_bytes())
        payload = backend("gui-run", "--op", "clean", "--opt", "scope.table=0",
                          "--opt", "scope.headers=0", "--opt", "rule.units=0",
                          "--files", str(selective))
        outputs = successful(payload, selective, root)
        selected = Document(next(path for path in outputs if path.suffix == ".docx"))
        assert selected.paragraphs[0].text == "正文：“成章”，面积10平方米，小时候。"
        assert selected.tables[0].cell(0, 0).text == table
        assert selected.sections[0].header.paragraphs[0].text == header
        assert sha256(selective) == original
        checked.append("GUI option forwarding: disabled table/header/unit rules, unspecified rules retain defaults")

        markdown = root / "引号示例.md"
        markdown.write_text('# 标题\n\n正文:"成章",面积10平方米。\n', encoding="utf-8")
        before = sha256(markdown)
        payload = backend("gui-run", "--op", "quotes", "--files", str(markdown))
        outputs = successful(payload, markdown, root)
        assert outputs[0].read_text(encoding="utf-8") == '# 标题\n\n正文:“成章”,面积10平方米。\n'
        assert sha256(markdown) == before
        checked.append("Markdown quotes: only quotation marks change; punctuation, units and source retained")

        first, second = root / "第一章.md", root / "第二章.md"
        first.write_text("# 第一章\n\n第一份内容。\n", encoding="utf-8")
        second.write_text("# 第二章\n\n第二份内容。\n", encoding="utf-8")
        hashes = (sha256(first), sha256(second))
        payload = backend("gui-run", "--op", "merge", "--files", str(first), str(second))
        outputs = successful(payload, first, root)
        merged = next(path for path in outputs if path.suffix == ".md").read_text(encoding="utf-8")
        assert "第一份内容。" in merged and "第二份内容。" in merged
        assert merged.index("第一份内容。") < merged.index("第二份内容。")
        assert (sha256(first), sha256(second)) == hashes
        checked.append("Markdown merge: ordered content, reported output and both sources unchanged")

        # Agent CLI: `dockit run` is the same gui_run; the exit code carries the outcome.
        agent = root / "agent 输入.docx"
        agent.write_bytes(source.read_bytes())
        code, payload = cli("run", "clean", "--opt", "scope.table=0", agent)
        assert code == 0, payload
        outputs = successful(payload, agent, root)
        assert Document(next(p for p in outputs if p.suffix == ".docx")).tables[0].cell(0, 0).text == table
        assert sha256(agent) == original
        listing = sorted(path.name for path in root.iterdir())
        for args, expected in ((["run", "convert", "--to", "bogus", agent], "bad_target"),
                               (["run", "renum", "--to", "bogus", agent], "bad_target"),
                               (["run", "stripchrome", agent], "confirm_required"),
                               (["run", "scan", root], "consent_required"),
                               # merged.md 已由上面的 GUI 合并写出:再合并不许静默覆盖
                               (["run", "merge", first, second], "would_overwrite")):
            code, payload = cli(*args)
            assert code == 2 and payload["error_code"] == expected, (args, code, payload)
            assert sorted(path.name for path in root.iterdir()) == listing, args
        code, payload = cli("run", "clean", agent, root / "absent.docx")
        assert code == 1 and payload["skipped_missing"] == ["absent.docx"], payload
        code, catalog_cli = cli("ops")
        assert code == 0 and catalog_cli == catalog
        checked.append("dockit CLI: run clean exit 0 with GUI envelope; invalid target, missing --yes, "
                       "scan without consent and overwriting an existing merged.md exit 2 before writing; "
                       "skipped input exit 1; ops equals gui-ops")

    finish("functionality", "真实 GUI 后端目录、4 条文件处理链及 dockit 命令行退出码通过", catalog_operations=len(ops),
           catalog_options=option_count, checks=checked,
           scope="Temporary DOCX/Markdown fixtures; no GUI activation, no network operation, no external document access")


if __name__ == "__main__":
    main()
